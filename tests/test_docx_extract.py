import os
import tempfile

import utils1
from docx import Document
from docx.oxml import OxmlElement


def test_table_and_paragraph_text_all_captured():
    doc = Document()
    doc.add_paragraph("Top Title")
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Voice over narration here"
    t.rows[0].cells[1].text = "Visual b-roll note"
    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    try:
        doc.save(path)
        out = utils1.load_script_file(path)
    finally:
        os.remove(path)
    assert "Top Title" in out
    assert "Voice over narration here" in out
    assert "Visual b-roll note" in out


def test_paragraph_captures_hyperlink_text():
    # python-docx's .runs skips <w:hyperlink> runs; our walker must still capture them.
    doc = Document()
    p = doc.add_paragraph("See: ")
    hyper = OxmlElement("w:hyperlink")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "https://drive.google.com/folder/xyz"
    r.append(t)
    hyper.append(r)
    p._p.append(hyper)
    text = utils1._paragraph_text_with_breaks(p)
    assert "https://drive.google.com/folder/xyz" in text
