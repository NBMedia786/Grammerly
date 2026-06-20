"""Two-column VO|Visuals extraction: structure detection + offset integrity."""
import os
import sys
import tempfile

from docx import Document

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from utils1 import load_script_structured


def _make_table_docx(rows, headers=("Voice Over", "Visuals")):
    doc = Document()
    t = doc.add_table(rows=0, cols=2)
    hr = t.add_row().cells
    hr[0].text, hr[1].text = headers
    for vo, vis in rows:
        c = t.add_row().cells
        c[0].text, c[1].text = vo, vis
    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    doc.save(path)
    return path


def test_two_column_layout_and_offsets():
    path = _make_table_docx([
        ("Intro. Officers arrive at the scene.", "bodycam_clip1.mp4"),
        ("They find the door open.", ""),
        ("The case takes a dark turn.", "https://drive.google.com/file/x"),
    ])
    try:
        r = load_script_structured(path)
    finally:
        os.remove(path)
    assert r["layout"] is not None
    rows, text = r["layout"]["rows"], r["text"]
    assert len(rows) == 3
    # the header row is skipped; VO text is the analyzed string
    assert "Officers arrive" in text and "dark turn" in text
    assert "Voice Over" not in text and "Visuals" not in text
    # offsets reconstruct the full text exactly (no gaps/overlaps)
    assert "\n\n".join(text[x["vo_start"]:x["vo_end"]] for x in rows) == text
    # visuals captured per row
    assert rows[0]["visuals"] == "bodycam_clip1.mp4"
    assert rows[1]["visuals"] == ""
    assert rows[2]["visuals"].startswith("https://drive.google.com")


def test_two_column_table_without_visuals_is_single_column():
    # A 2-col table where the Visuals column is entirely empty -> treat as single column.
    path = _make_table_docx([("Some narration line one.", ""), ("Another narration line.", "")])
    try:
        r = load_script_structured(path)
    finally:
        os.remove(path)
    assert r["layout"] is None
    assert "narration line one" in r["text"]


def test_plain_paragraph_doc_has_no_layout():
    doc = Document()
    doc.add_paragraph("This is a normal single-column paragraph script with plenty of words.")
    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    doc.save(path)
    try:
        r = load_script_structured(path)
    finally:
        os.remove(path)
    assert r["layout"] is None
    assert "single-column paragraph" in r["text"]
